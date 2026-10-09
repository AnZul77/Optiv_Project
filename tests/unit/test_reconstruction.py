"""
Unit Tests for Document Reconstruction & Integrity Verification
Person 1: Lead Architect
"""

import os
import io
import pytest
from docx import Document
from pptx import Presentation
import fitz

from src.security.integrity import DocumentIntegrityGuard
from src.sanitization.docx_reconstructor import DOCXReconstructor


@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "original.docx"
    doc = Document()
    doc.add_heading("Confidential TPRM Audit", level=1)
    doc.add_paragraph("Employee Name: Alice Vance, SSN: 987-65-4321.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Vendor"
    table.rows[0].cells[1].text = "Contact Email"
    table.rows[1].cells[0].text = "Acme Corp"
    table.rows[1].cells[1].text = "alice@acme.com"

    doc.save(str(docx_path))
    return str(docx_path)


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = tmp_path / "test.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Sample text", fontsize=12)
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


def test_document_integrity_guard_valid(sample_docx, sample_pdf):
    # Test valid DOCX
    res_docx = DocumentIntegrityGuard.verify_document(sample_docx)
    assert res_docx.is_valid is True
    assert res_docx.file_type == ".docx"

    # Test valid PDF
    res_pdf = DocumentIntegrityGuard.verify_document(sample_pdf)
    assert res_pdf.is_valid is True
    assert res_pdf.file_type == ".pdf"


def test_document_integrity_guard_corrupt(tmp_path):
    # Test corrupted DOCX (random bytes)
    bad_docx = tmp_path / "corrupt.docx"
    bad_docx.write_bytes(b"PK\x03\x04randomjunknotazip")
    
    res = DocumentIntegrityGuard.verify_document(str(bad_docx))
    assert res.is_valid is False
    assert "Corrupted" in res.error or "failed" in res.error

    # Test missing file
    res_missing = DocumentIntegrityGuard.verify_document(str(tmp_path / "nonexistent.docx"))
    assert res_missing.is_valid is False
    assert "File not found" in res_missing.error


def test_docx_reconstructor_text_replacement(sample_docx, tmp_path):
    output_path = tmp_path / "reconstructed.docx"
    replacements = {
        "Alice Vance": "[REDACTED_PERSON]",
        "987-65-4321": "[REDACTED_SSN]",
        "alice@acme.com": "[REDACTED_EMAIL]",
    }

    success, err = DOCXReconstructor.replace_text_in_document(
        sample_docx,
        str(output_path),
        replacements
    )

    assert success is True
    assert err is None
    assert os.path.exists(str(output_path))

    # Verify content in reconstructed file
    reconstructed_doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in reconstructed_doc.paragraphs)
    
    assert "[REDACTED_PERSON]" in full_text
    assert "[REDACTED_SSN]" in full_text
    assert "Alice Vance" not in full_text
    assert "987-65-4321" not in full_text

    # Verify table replacement
    table_text = "\n".join(cell.text for row in reconstructed_doc.tables[0].rows for cell in row.cells)
    assert "[REDACTED_EMAIL]" in table_text
    assert "alice@acme.com" not in table_text


def test_docx_reconstructor_image_replacement(tmp_path):
    # Create docx with a tiny embedded image
    img_docx_path = tmp_path / "doc_with_img.docx"
    doc = Document()
    doc.add_paragraph("Document with embedded image:")
    
    # Create 1x1 png image
    dummy_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    img_stream = io.BytesIO(dummy_png)
    doc.add_picture(img_stream)
    doc.save(str(img_docx_path))

    # New burned pixel dummy bytes
    new_burned_bytes = dummy_png  # valid png replacement

    out_docx_path = tmp_path / "doc_redacted_img.docx"
    
    # Get image relation ID
    rel_ids = [rel_id for rel_id, rel in doc.part.rels.items() if "image" in rel.target_ref.lower()]
    assert len(rel_ids) > 0
    target_rel = rel_ids[0]

    success, err = DOCXReconstructor.replace_images_in_document(
        str(img_docx_path),
        str(out_docx_path),
        {target_rel: new_burned_bytes}
    )

    assert success is True
    assert err is None
    assert os.path.exists(str(out_docx_path))
    
    # Ensure integrity check passes
    check = DocumentIntegrityGuard.verify_document(str(out_docx_path))
    assert check.is_valid is True
