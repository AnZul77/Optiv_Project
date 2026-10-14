"""
Unit and integration tests for DOCX embedded image OCR and redaction.
Tests extraction, detection, pixel burning, and reconstruction integrity.
"""

import io
import os
from pathlib import Path
from typing import Any, Dict

import numpy as np
import pytest
from docx import Document
from PIL import Image, ImageDraw

from src.sanitization.docx_image_redactor import (
    DOCXImageRedactor,
    EmbeddedImageInfo,
    DOCXImageRedactionReport,
)
from src.sanitization.docx_reconstructor import DOCXReconstructor
from src.security.integrity import DocumentIntegrityGuard


def _create_synthetic_image(
    text: str = "",
    width: int = 200,
    height: int = 80,
    bg_color: tuple = (255, 255, 255),
    text_color: tuple = (0, 0, 0),
    fmt: str = "PNG",
) -> bytes:
    """Helper to generate an in-memory image with optional text."""
    img = Image.new("RGB", (width, height), color=bg_color)
    if text:
        draw = ImageDraw.Draw(img)
        draw.text((10, 25), text, fill=text_color)
    buf = io.BytesIO()
    img.save(buf, format=fmt)
    return buf.getvalue()


def _create_sample_docx(docx_path: str | Path, images: list[bytes]) -> None:
    """Helper to create a DOCX file containing paragraphs and embedded images."""
    doc = Document()
    doc.add_heading("Confidential TPRM Training Manual", level=1)
    doc.add_paragraph("The following screenshot demonstrates system access:")

    for idx, img_bytes in enumerate(images):
        stream = io.BytesIO(img_bytes)
        doc.add_picture(stream)
        doc.add_paragraph(f"Figure {idx + 1}: Embedded System Interface")

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(str(docx_path))


class TestDOCXImageExtraction:
    """Test extracting embedded images from DOCX documents."""

    def test_extract_single_and_multiple_images(self, tmp_path):
        docx_file = tmp_path / "test_extract.docx"
        img1 = _create_synthetic_image("User: admin", 150, 60)
        img2 = _create_synthetic_image("Role: Manager", 180, 70)
        _create_sample_docx(docx_file, [img1, img2])

        redactor = DOCXImageRedactor()
        images = redactor.extract_images(str(docx_file))

        assert len(images) == 2
        for img_info in images:
            assert isinstance(img_info, EmbeddedImageInfo)
            assert img_info.is_raster is True
            assert img_info.width is not None and img_info.width > 0
            assert img_info.height is not None and img_info.height > 0
            assert len(img_info.blob) > 0
            assert img_info.filename.endswith(".png")

    def test_extract_from_open_document_object(self, tmp_path):
        docx_file = tmp_path / "test_open_doc.docx"
        img = _create_synthetic_image("Test doc", 100, 50)
        _create_sample_docx(docx_file, [img])

        doc = Document(str(docx_file))
        redactor = DOCXImageRedactor()
        images = redactor.extract_images(doc)

        assert len(images) == 1
        assert images[0].width == 100
        assert images[0].height == 50


class TestPIIDetectionAndImageRedaction:
    """Test PII detection on OCR items and pixel burning."""

    def test_detect_pii_boxes_with_pre_extracted_ocr(self):
        redactor = DOCXImageRedactor()
        img_bytes = _create_synthetic_image("Employee EMP-908234", 300, 100)

        mock_ocr = [
            {"text": "Employee EMP-908234", "pixel_bbox": [10, 20, 250, 60]}
        ]
        boxes = redactor.detect_pii_boxes(img_bytes, ocr_results=mock_ocr)

        assert len(boxes) >= 1
        x0, y0, x1, y1 = boxes[0]
        assert 0.0 <= x0 < x1 <= 1.0
        assert 0.0 <= y0 < y1 <= 1.0

    def test_detect_pii_boxes_clean_image_returns_empty(self):
        redactor = DOCXImageRedactor()
        img_bytes = _create_synthetic_image("System Diagram Architecture", 200, 80)

        mock_ocr = [
            {"text": "System Diagram Architecture", "pixel_bbox": [10, 10, 190, 70]}
        ]
        boxes = redactor.detect_pii_boxes(img_bytes, ocr_results=mock_ocr)
        assert boxes == []

    def test_redact_image_burns_black_pixels(self):
        redactor = DOCXImageRedactor(padding=0)
        # 100x100 white image
        img_bytes = _create_synthetic_image(width=100, height=100)

        # Target normalized box in the center: [0.3, 0.3, 0.7, 0.7] (pixels 30..70)
        box = [0.3, 0.3, 0.7, 0.7]
        sanitized_bytes = redactor.redact_image(img_bytes, [box])

        # Verify pixels
        with Image.open(io.BytesIO(sanitized_bytes)) as out_img:
            arr = np.array(out_img)
            # Center must be solid black (0, 0, 0)
            assert np.all(arr[40:60, 40:60] == 0)
            # Corner must remain white (255, 255, 255)
            assert np.all(arr[5, 5] == 255)


class TestDOCXEndToEndRedaction:
    """Test full document image replacement and post-reconstruction integrity."""

    def test_end_to_end_image_redaction_and_replacement(self, tmp_path):
        in_docx = tmp_path / "source_with_pii_image.docx"
        out_docx = tmp_path / "redacted_output.docx"

        # Create image with PII text
        img_bytes = _create_synthetic_image("SSN: 900-12-3456", 200, 60)
        _create_sample_docx(in_docx, [img_bytes])

        # Mock OCR identifying the PII
        pre_extracted = {
            "image1.png": [
                {"text": "SSN: 900-12-3456", "pixel_bbox": [10, 15, 180, 45]}
            ]
        }

        redactor = DOCXImageRedactor(padding=2)
        success, err, report = redactor.redact_docx(
            input_path=in_docx,
            output_path=out_docx,
            pre_extracted_ocr=pre_extracted,
            preserve_integrity=True,
        )

        assert success is True
        assert err is None
        assert report.total_images == 1
        assert report.redacted_images == 1
        assert report.total_boxes_redacted >= 1
        assert os.path.exists(str(out_docx))

        # 1. Verify DocumentIntegrityGuard validates reconstructed file
        integrity = DocumentIntegrityGuard.verify_document(str(out_docx))
        assert integrity.is_valid is True

        # 2. Re-open output docx, inspect extracted image, verify black pixels
        out_images = redactor.extract_images(str(out_docx))
        assert len(out_images) == 1
        with Image.open(io.BytesIO(out_images[0].blob)) as reloaded_img:
            arr = np.array(reloaded_img)
            # Check center of the redacted SSN region (x=90..140, y=25..35) is black
            assert np.all(arr[25:35, 90:140] == 0)

    def test_document_with_no_pii_images_passes_unchanged(self, tmp_path):
        in_docx = tmp_path / "clean_images.docx"
        out_docx = tmp_path / "clean_output.docx"

        img_bytes = _create_synthetic_image("Overview Chart", 120, 50)
        _create_sample_docx(in_docx, [img_bytes])

        pre_extracted = {
            "image1.png": [
                {"text": "Overview Chart", "pixel_bbox": [10, 10, 100, 40]}
            ]
        }

        redactor = DOCXImageRedactor()
        success, err, report = redactor.redact_docx(
            input_path=in_docx,
            output_path=out_docx,
            pre_extracted_ocr=pre_extracted,
        )

        assert success is True
        assert err is None
        assert report.total_images == 1
        assert report.redacted_images == 0
        assert report.total_boxes_redacted == 0
        assert os.path.exists(str(out_docx))

        integrity = DocumentIntegrityGuard.verify_document(str(out_docx))
        assert integrity.is_valid is True

    def test_docx_reconstructor_delegation(self, tmp_path):
        in_docx = tmp_path / "delegate_in.docx"
        out_docx = tmp_path / "delegate_out.docx"

        img_bytes = _create_synthetic_image("Contact: priyanka@cadencefg.com", 250, 70)
        _create_sample_docx(in_docx, [img_bytes])

        pre_extracted = {
            "image1.png": [
                {"text": "Contact: priyanka@cadencefg.com", "pixel_bbox": [10, 20, 230, 50]}
            ]
        }

        # Call via DOCXReconstructor.redact_images_in_document
        success, err = DOCXReconstructor.redact_images_in_document(
            input_path=str(in_docx),
            output_path=str(out_docx),
            pre_extracted_ocr=pre_extracted,
        )

        assert success is True
        assert err is None
        assert os.path.exists(str(out_docx))
        assert DocumentIntegrityGuard.verify_document(str(out_docx)).is_valid is True

    def test_custom_mock_ocr_engine_object(self, tmp_path):
        in_docx = tmp_path / "custom_engine_in.docx"
        out_docx = tmp_path / "custom_engine_out.docx"

        img_bytes = _create_synthetic_image("Delphine Okonjo", 200, 60)
        _create_sample_docx(in_docx, [img_bytes])

        class MockOCREngine:
            def extract(self, image_data):
                return {
                    "words": [
                        {
                            "word": "Delphine Okonjo",
                            "pixel_bbox": [10, 15, 180, 45],
                            "confidence": 0.95,
                        }
                    ]
                }

        redactor = DOCXImageRedactor(ocr_engine=MockOCREngine())
        success, err, report = redactor.redact_docx(
            input_path=in_docx,
            output_path=out_docx,
        )

        assert success is True
        assert report.redacted_images == 1
        assert report.total_boxes_redacted >= 1
        assert DocumentIntegrityGuard.verify_document(str(out_docx)).is_valid is True

    def test_nonexistent_source_fails_gracefully(self, tmp_path):
        redactor = DOCXImageRedactor()
        success, err, report = redactor.redact_docx(
            input_path=tmp_path / "nonexistent.docx",
            output_path=tmp_path / "output.docx",
        )
        assert success is False
        assert "not found" in str(err).lower()
