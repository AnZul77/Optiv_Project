import pytest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.sanitization.pptx_reconstructor import PPTXReconstructor


def test_pptx_reconstructor_shapes_tables_notes(tmp_path):
    source_pptx = tmp_path / "test_presentation.pptx"
    output_pptx = tmp_path / "sanitized_presentation.pptx"

    # Create a synthetic presentation
    prs = Presentation()
    prs.core_properties.author = "Sensitive Author"
    prs.core_properties.title = "Internal Strategy"

    # Slide 1: Shape text
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    tx_box = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf = tx_box.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "CEO: Marguerite Vossberg"
    p2 = tf.add_paragraph()
    p2.text = "Contact: m.vossberg@cadencefg.example"

    # Slide 2: Table
    slide2 = prs.slides.add_slide(blank_layout)
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Holder"
    table.cell(1, 0).text = "Chief Information Officer"
    table.cell(1, 1).text = "Tobias Lindqvist"

    # Slide 2 Notes
    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = "Private note: Review with Tobias Lindqvist next Tuesday."

    prs.save(str(source_pptx))

    replacements = {
        "Marguerite Vossberg": "[REDACTED_PERSON]",
        "m.vossberg@cadencefg.example": "[REDACTED_EMAIL]",
        "Tobias Lindqvist": "[REDACTED_PERSON]",
    }

    success, err = PPTXReconstructor.replace_text_in_presentation(
        source_pptx,
        output_pptx,
        replacements
    )

    assert success is True
    assert err is None
    assert output_pptx.exists()

    # Verify sanitized output
    sanitized_prs = Presentation(str(output_pptx))
    assert sanitized_prs.core_properties.author == ""
    assert sanitized_prs.core_properties.title == ""

    # Check Slide 1
    s1_text = []
    for shape in sanitized_prs.slides[0].shapes:
        if shape.has_text_frame:
            s1_text.extend([p.text for p in shape.text_frame.paragraphs])
    assert "[REDACTED_PERSON]" in s1_text[0]
    assert "Marguerite Vossberg" not in s1_text[0]
    assert "[REDACTED_EMAIL]" in s1_text[1]
    assert "m.vossberg@cadencefg.example" not in s1_text[1]

    # Check Slide 2 Table
    s2_table = sanitized_prs.slides[1].shapes[0].table
    assert s2_table.cell(0, 0).text == "Role"
    assert s2_table.cell(1, 0).text == "Chief Information Officer"
    assert s2_table.cell(1, 1).text == "[REDACTED_PERSON]"
    assert "Tobias Lindqvist" not in s2_table.cell(1, 1).text

    # Check Slide 2 Notes
    notes_text = sanitized_prs.slides[1].notes_slide.notes_text_frame.text
    assert "[REDACTED_PERSON]" in notes_text
    assert "Tobias Lindqvist" not in notes_text
