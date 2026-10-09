"""
Unit Tests for Document Ingestion and Dispatcher
Person 1: Lead Architect
"""

import os
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
import fitz

from src.ingestion.validators import validate_document, inspect_magic_bytes
from src.ingestion.pdf import PDFExtractor
from src.ingestion.docx import DOCXExtractor
from src.ingestion.pptx import PPTXExtractor
from src.ingestion.metadata import MetadataExtractor, MetadataSanitizer
from src.ingestion.dispatcher import dispatch_extraction, route_document, list_supported_formats
from src.schema.document import CanonicalDocument, ExtractedPage, ExtractedBlock, EntityAnnotation


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "CONFIDENTIAL RISK REPORT\nOfficer: Alice Smith\nSSN: 000-12-3456", fontsize=12)
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)


@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("TPRM Assessment", level=1)
    doc.add_paragraph("Vendor primary contact details below.")
    
    # Add a table with headers and data
    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Vendor Name"
    hdr_cells[1].text = "Contact Email"
    hdr_cells[2].text = "Tax ID"
    
    row_cells = table.rows[1].cells
    row_cells[0].text = "Acme Corp"
    row_cells[1].text = "contact@acme.com"
    row_cells[2].text = "XX-1234567"
    
    doc.core_properties.author = "John Examiner"
    doc.save(str(docx_path))
    return str(docx_path)


@pytest.fixture
def sample_pptx(tmp_path):
    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Executive Leadership Org Pack"
    
    # Add speaker note
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Secret note: CEO direct phone +1-555-0199"
        
    prs.core_properties.author = "HR Director"
    prs.save(str(pptx_path))
    return str(pptx_path)


def test_validate_document_success(sample_pdf, sample_docx, sample_pptx):
    res_pdf = validate_document(sample_pdf)
    assert res_pdf.is_valid is True
    assert res_pdf.detected_type == ".pdf"

    res_docx = validate_document(sample_docx)
    assert res_docx.is_valid is True
    assert res_docx.detected_type == ".docx"

    res_pptx = validate_document(sample_pptx)
    assert res_pptx.is_valid is True
    assert res_pptx.detected_type == ".pptx"


def test_validate_document_mismatch(tmp_path):
    # Fake file claiming to be PDF
    fake_pdf = tmp_path / "fake.pdf"
    fake_pdf.write_text("This is plain text, not a PDF.")
    
    res = validate_document(str(fake_pdf))
    assert res.is_valid is False
    assert "Magic byte mismatch" in res.error


def test_pdf_extraction(sample_pdf):
    extractor = PDFExtractor()
    doc, meta = extractor.extract(sample_pdf)

    assert doc.pages == 1
    assert doc.file_type == ".pdf"
    assert "Alice Smith" in doc.get_full_text()
    assert len(doc.pages_dict[1].blocks) > 0
    assert meta.file_size_bytes > 0


def test_docx_extraction(sample_docx):
    extractor = DOCXExtractor()
    doc, meta = extractor.extract(sample_docx)

    assert doc.file_type == ".docx"
    assert len(doc.get_all_tables()) == 1
    
    # Verify table column header inheritance
    table = doc.get_all_tables()[0]
    assert table.headers == ["Vendor Name", "Contact Email", "Tax ID"]
    assert table.rows[0][1].header_name == "Contact Email"
    assert table.rows[0][1].text == "contact@acme.com"


def test_pptx_extraction(sample_pptx):
    extractor = PPTXExtractor()
    doc, meta = extractor.extract(sample_pptx)

    assert doc.file_type == ".pptx"
    assert "Executive Leadership" in doc.get_full_text()
    # Check speaker notes extraction
    notes_blocks = [b for b in doc.pages_dict[1].blocks if b.block_type == "speaker_note"]
    assert len(notes_blocks) == 1
    assert "+1-555-0199" in notes_blocks[0].text


def test_dispatcher_routing(sample_pdf, sample_docx, sample_pptx):
    formats = list_supported_formats()
    assert ".pdf" in formats
    assert ".docx" in formats
    assert ".pptx" in formats

    res = dispatch_extraction(sample_pdf)
    assert res.success is True
    assert res.document is not None

    route_summary = route_document(sample_docx)
    assert route_summary["status"] == "EXTRACTED"
    assert route_summary["total_tables"] == 1


def test_metadata_extraction_and_sanitization(sample_docx, tmp_path):
    meta = MetadataExtractor.extract_from_docx(sample_docx)
    assert meta["author"] == "John Examiner"

    sanitized_docx = tmp_path / "sanitized.docx"
    ok = MetadataSanitizer.sanitize_docx_properties(sample_docx, str(sanitized_docx))
    assert ok is True

    clean_meta = MetadataExtractor.extract_from_docx(str(sanitized_docx))
    assert clean_meta["author"] == "[REDACTED]"


def test_canonical_document_entity_creation():
    doc = CanonicalDocument(
        doc_id="test_doc",
        source_path="fake.pdf",
        filename="fake.pdf",
        file_type=".pdf",
        pages=1
    )
    
    entity = EntityAnnotation.create(
        entity_type="SSN",
        raw_value="123-45-6789",
        source="regex",
        page=1,
        bbox=[10.0, 20.0, 50.0, 30.0],
        risk="CRITICAL",
        action="REDACT"
    )
    
    assert entity.type == "SSN"
    assert len(entity.value_hash) == 64  # SHA-256 hash length
    assert "123-45-6789" not in entity.value_hash  # raw PII strictly omitted
    assert entity.risk == "CRITICAL"
