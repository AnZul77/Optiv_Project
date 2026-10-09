"""
PPTX Extractor
Person 1: Lead Architect
Handles PPTX extraction including slide shapes, group shapes/diagrams (org charts),
speaker notes (hidden PII channels), hidden slides (<p:sld show="0">), and slide tables.
Roadmap: Day 4 - PPTX Processing & NER Integration
"""

from __future__ import annotations
import os
import time
from typing import Dict, List, Optional, Tuple, Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..schema.document import (
    CanonicalDocument,
    ExtractedPage,
    ExtractedBlock,
    ExtractedImage,
    ExtractedTable,
    TableCell,
    PipelineMetadata,
    ExtractionResult,
)


class PPTXExtractor:
    """
    Presentation extraction handler that supports:
    - Native slide shapes & hierarchical group shapes (org charts)
    - Hidden slides detection via OpenXML show="0" attribute
    - Speaker notes extraction (critical covert PII channel)
    - Embedded diagrams, tables, and images
    """

    def __init__(self, ocr_engine=None):
        """
        Args:
            ocr_engine: Optional OCR engine adapter (from Person 2).
        """
        self.ocr_engine = ocr_engine

    def extract(
        self,
        file_path: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> Tuple[CanonicalDocument, PipelineMetadata]:
        """
        Extract content from a PPTX file.
        
        Args:
            file_path: Path to the PPTX file.
            metadata: Additional metadata dict.
            
        Returns:
            Tuple of (CanonicalDocument, PipelineMetadata).
        """
        start_time = time.time()
        file_meta = metadata or {}
        doc_id = file_meta.get("doc_id", os.path.basename(file_path))
        filename = os.path.basename(file_path)
        file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0

        pipeline_meta = PipelineMetadata(
            file_size_bytes=file_size,
            extractor_version="1.0.0-pptx",
            ocr_engine="python-pptx"
        )

        prs = Presentation(file_path)
        total_slides = len(prs.slides)

        canonical_doc = CanonicalDocument(
            doc_id=doc_id,
            source_path=file_path,
            filename=filename,
            file_type=".pptx",
            pages=total_slides,
            metadata={
                "slide_count": total_slides,
                "author": prs.core_properties.author or "",
                "title": prs.core_properties.title or "",
            }
        )

        hidden_slides_count = 0
        notes_detected_count = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            extracted_page = ExtractedPage(page_num=slide_idx)
            
            # 1. Detect if slide is marked hidden in OpenXML
            # In PPTX XML: <p:sld show="0"> indicates a hidden slide
            is_slide_hidden = (slide._element.get("show") == "0")
            if is_slide_hidden:
                hidden_slides_count += 1
                extracted_page.metadata["is_hidden_slide"] = True
                pipeline_meta.warnings.append(f"Slide {slide_idx} is marked as HIDDEN")

            # 2. Extract shapes recursively (including group shapes / org charts)
            slide_blocks: List[ExtractedBlock] = []
            slide_images: List[ExtractedImage] = []
            slide_tables: List[ExtractedTable] = []

            self._process_shapes(
                slide.shapes,
                slide_idx=slide_idx,
                blocks=slide_blocks,
                images=slide_images,
                tables=slide_tables
            )

            # 3. Extract speaker notes (hidden PII channel)
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide and notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        notes_detected_count += 1
                        slide_blocks.append(ExtractedBlock(
                            block_id=f"pptx_s{slide_idx}_notes",
                            page=slide_idx,
                            block_type="speaker_note",
                            text=notes_text,
                            is_hidden=True,
                            metadata={"is_speaker_note": True}
                        ))

            # Assemble page text
            extracted_page.blocks = slide_blocks
            extracted_page.images = slide_images
            extracted_page.tables = slide_tables
            
            native_text_parts = [b.text for b in slide_blocks if b.text.strip()]
            extracted_page.native_text = "\n\n".join(native_text_parts)
            extracted_page.has_native_text = bool(extracted_page.native_text)
            extracted_page.ocr_confidence = 1.0

            canonical_doc.add_page(extracted_page)

        canonical_doc.metadata["hidden_slides_count"] = hidden_slides_count
        canonical_doc.metadata["notes_count"] = notes_detected_count
        
        if hidden_slides_count > 0:
            pipeline_meta.notes.append(f"{hidden_slides_count} hidden slide(s) identified")
        if notes_detected_count > 0:
            pipeline_meta.notes.append(f"Speaker notes detected across {notes_detected_count} slide(s)")

        pipeline_meta.extraction_time = time.time() - start_time
        return canonical_doc, pipeline_meta

    def _process_shapes(
        self,
        shapes,
        slide_idx: int,
        blocks: List[ExtractedBlock],
        images: List[ExtractedImage],
        tables: List[ExtractedTable]
    ):
        """Recursively process shapes, nested groups, tables, and pictures."""
        for shape_idx, shape in enumerate(shapes):
            # Check for GroupShape (hierarchical diagrams & org charts)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                self._process_shapes(
                    shape.shapes,
                    slide_idx=slide_idx,
                    blocks=blocks,
                    images=images,
                    tables=tables
                )
                continue

            # Check for Table Shape
            if shape.has_table:
                table = shape.table
                all_rows = list(table.rows)
                if all_rows:
                    headers = [cell.text.strip() for cell in all_rows[0].cells]
                    table_rows = []
                    for r_idx, row in enumerate(all_rows[1:], start=1):
                        row_cells = []
                        for c_idx, cell in enumerate(row.cells):
                            h_name = headers[c_idx] if c_idx < len(headers) else f"Col_{c_idx+1}"
                            row_cells.append(TableCell(
                                row_idx=r_idx,
                                col_idx=c_idx,
                                text=cell.text.strip(),
                                header_name=h_name
                            ))
                        table_rows.append(row_cells)
                    
                    ext_tbl = ExtractedTable(
                        table_index=len(tables) + 1,
                        page=slide_idx,
                        headers=headers,
                        rows=table_rows
                    )
                    ext_tbl.raw_markdown = ext_tbl.to_markdown()
                    tables.append(ext_tbl)
                continue

            # Check for Text Frame
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    blocks.append(ExtractedBlock(
                        block_id=f"pptx_s{slide_idx}_sh_{shape_idx}",
                        page=slide_idx,
                        block_type="slide_shape",
                        text=txt,
                        metadata={"shape_name": shape.name}
                    ))

            # Check for Picture / Image shape
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    img = shape.image
                    img_bytes = img.blob
                    ext = img.ext
                    images.append(ExtractedImage(
                        image_id=f"pptx_s{slide_idx}_img_{shape_idx}",
                        page=slide_idx,
                        image_bytes=img_bytes,
                        image_format=ext,
                        width=img.size[0] if hasattr(img, "size") else 0,
                        height=img.size[1] if hasattr(img, "size") else 0,
                        metadata={"shape_name": shape.name}
                    ))
                except Exception:
                    pass