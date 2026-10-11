"""
Native PPTX Reconstructor
Person 1 & 2: Multimodal Document Sanitization
Roadmap: Reconstruct sanitized PPTX presentations with layout and style retention.

Replaces sensitive text within shapes, group shapes (hierarchies/org charts),
tables, and speaker notes while preserving PowerPoint themes, fonts, geometries,
and visual structure. Clears core document metadata.
"""

from __future__ import annotations
import os
import shutil
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .text_redactor import redact_text_by_values


class PPTXReconstructor:
    """
    Reconstructs PowerPoint presentations with redacted text in-place.
    Preserves presentation geometries, layout templates, master slides,
    group shapes, and formatting while replacing sensitive entities.
    """

    @classmethod
    def replace_text_in_presentation(
        cls,
        source_path: str | Path,
        destination_path: str | Path,
        replacements: Dict[str, str],
    ) -> Tuple[bool, Optional[str]]:
        """
        Replaces sensitive text strings in a presentation with their redaction tokens.

        Args:
            source_path: Path to original .pptx file.
            destination_path: Path to output .pptx file.
            replacements: Dict mapping sensitive strings -> replacement tokens.

        Returns:
            Tuple of (success: bool, error_message: Optional[str])
        """
        source = Path(source_path)
        dest = Path(destination_path)

        if not source.exists():
            return False, f"Source PPTX file not found: {source}"

        dest.parent.mkdir(parents=True, exist_ok=True)

        try:
            prs = Presentation(str(source))

            for slide in prs.slides:
                cls._process_shapes(slide.shapes, replacements)

                # Process speaker notes if present (hidden PII channel)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    cls._redact_text_frame(slide.notes_slide.notes_text_frame, replacements)

            # Clear core document metadata
            try:
                prs.core_properties.author = ""
                prs.core_properties.last_modified_by = ""
                prs.core_properties.title = ""
                prs.core_properties.comments = ""
            except Exception:
                pass

            prs.save(str(dest))
            return True, None

        except Exception as e:
            return False, str(e)

    @classmethod
    def _process_shapes(cls, shapes, replacements: Dict[str, str]):
        """Recursively process shapes, nested groups, and tables."""
        for shape in shapes:
            # 1. Group shapes (org charts, nested diagram containers)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                cls._process_shapes(shape.shapes, replacements)
                continue

            # 2. Slide tables
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            cls._redact_text_frame(cell.text_frame, replacements)
                continue

            # 3. Text frames (shapes, text boxes, callouts)
            if shape.has_text_frame:
                cls._redact_text_frame(shape.text_frame, replacements)

    @classmethod
    def _redact_text_frame(cls, text_frame, replacements: Dict[str, str]):
        """Redacts text within a text frame paragraph-by-paragraph."""
        for paragraph in text_frame.paragraphs:
            original_text = paragraph.text
            if not original_text or not any(k in original_text for k in replacements):
                continue

            redacted_text = redact_text_by_values(original_text, replacements)
            if redacted_text != original_text:
                if paragraph.runs:
                    paragraph.runs[0].text = redacted_text
                    for r in paragraph.runs[1:]:
                        r.text = ""
                else:
                    paragraph.text = redacted_text
